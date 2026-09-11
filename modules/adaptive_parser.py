from __future__ import annotations

import csv
import hashlib
import math
import re
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from .rab_models import ColumnMapping, MappingDiagnostic, ParsedDocument, RABItem, TableRegion, confidence_label


ROLE_ORDER = [
    "item_number",
    "description",
    "unit",
    "volume",
    "material_unit_price",
    "service_unit_price",
    "unit_price",
    "material_total",
    "service_total",
    "total_price",
    "notes",
]
NUMERIC_ROLES = {
    "volume",
    "material_unit_price",
    "service_unit_price",
    "unit_price",
    "material_total",
    "service_total",
    "total_price",
}
MARKERS = {"", "-", "--", "n/a", "na", "include", "included", "inklusif", "nihil", "none"}
_MERGED_VALUE_CACHE: dict[tuple[int, int, int], Any] = {}


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().lower()
    if text in MARKERS or text.startswith("="):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"(?i)\b(?:rp|idr)\b", "", text)
    if re.search(r"[a-z]", text, re.IGNORECASE):
        return None
    text = text.strip().strip("()")
    text = re.sub(r"[^0-9,\.\-+]", "", text)
    if not text or text in {"-", "+", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        text = "".join(parts) if len(parts[-1]) == 3 and len(parts) > 1 else ".".join(parts)
    elif "." in text:
        parts = text.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[-1]) == 3 and len(parts[0]) >= 1):
            text = "".join(parts)
    try:
        number = float(text)
    except ValueError:
        return None
    if negative:
        number = -abs(number)
    return number if math.isfinite(number) else None


def parse_document(
    file_path: str | Path,
    manual_mappings: dict[str, dict[str, Any]] | None = None,
) -> ParsedDocument:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(path, manual_mappings or {})
    if suffix == ".xls":
        return _parse_legacy_excel(path, manual_mappings or {})
    if suffix == ".csv":
        return _parse_csv(path, manual_mappings or {})
    raise ValueError(f"Parser tabel tidak mendukung format {suffix}.")


def parse_dataframe_tables(
    file_path: str | Path,
    tables: dict[str, pd.DataFrame],
    strategy: str,
    manual_mappings: dict[str, dict[str, Any]] | None = None,
    source_hash: str | None = None,
) -> ParsedDocument:
    path = Path(file_path)
    return _parse_dataframe_sheets(
        path,
        source_hash or file_sha256(path),
        tables,
        strategy,
        manual_mappings or {},
    )


def preview_dataframe(document: ParsedDocument) -> pd.DataFrame:
    rows = []
    for item in document.items:
        rows.append(
            {
                "No.": item.item_no or item.display_sequence,
                "Section": item.section,
                "Item / Uraian": item.item_per_rab,
                "Satuan": item.unit,
                "Volume": item.volume,
                "Harga Satuan": item.unit_price,
                "Total": item.total_price,
                "Sheet / Page": item.page_or_sheet,
                "Source Row": item.source_coordinate,
            }
        )
    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    return frame.dropna(axis=1, how="all").loc[:, lambda df: (df.astype(str) != "").any(axis=0)]


def _parse_xlsx(path: Path, manual_mappings: dict[str, dict[str, Any]]) -> ParsedDocument:
    from openpyxl import load_workbook

    source_hash = file_sha256(path)
    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    document = ParsedDocument(path.name, source_hash, "adaptive_xlsx")
    display_sequence = 0
    for formula_sheet in formula_book.worksheets:
        value_sheet = value_book[formula_sheet.title]
        meaningful_last_row = _meaningful_last_row(value_sheet, formula_sheet)
        if meaningful_last_row <= 0:
            continue
        title = _extract_sheet_title(value_sheet, meaningful_last_row) or path.stem
        candidates = _find_header_regions(formula_sheet, value_sheet, meaningful_last_row)
        if not candidates:
            document.warnings.append(f"{formula_sheet.title}: header tabel tidak ditemukan.")
            continue
        for region_index, candidate in enumerate(candidates, start=1):
            header_start, header_end = candidate
            region_id = f"{formula_sheet.title}:{header_start}-{header_end}"
            next_start = candidates[region_index][0] if region_index < len(candidates) else meaningful_last_row + 1
            data_end = max(header_end, next_start - 1)
            mapping = _map_columns(
                formula_sheet,
                value_sheet,
                header_start,
                header_end,
                data_end,
                manual_mappings.get(region_id, {}),
            )
            hidden_count = sum(
                1 for row_no in range(header_end + 1, data_end + 1) if formula_sheet.row_dimensions[row_no].hidden
            )
            confidence = _mapping_confidence(mapping)
            warnings = _mapping_warnings(mapping)
            review_blocked = _mapping_blocks_review(mapping, confidence)
            diagnostic = MappingDiagnostic(
                region_id=region_id,
                confidence=confidence,
                confidence_label=confidence_label(confidence),
                header_rows=list(range(header_start, header_end + 1)),
                mappings=mapping,
                hidden_rows_skipped=hidden_count,
                warnings=warnings,
                review_blocked=review_blocked,
            )
            items, arithmetic_failures, display_sequence = _extract_region_items(
                path,
                source_hash,
                formula_sheet,
                value_sheet,
                title,
                header_end + 1,
                data_end,
                mapping,
                confidence,
                display_sequence,
            )
            diagnostic.arithmetic_failures = arithmetic_failures
            if arithmetic_failures:
                diagnostic.warnings.append(f"{arithmetic_failures} baris memiliki konflik aritmetika.")
            if not items:
                diagnostic.warnings.append("Region tidak menghasilkan item RAB.")
            region = TableRegion(
                region_id,
                formula_sheet.title,
                header_start,
                header_end,
                header_end + 1,
                data_end,
                diagnostic,
            )
            document.regions.append(region)
            document.items.extend(items)
            if items and review_blocked:
                document.review_blocked = True
    if not document.items:
        document.warnings.append("Tidak ada item RAB yang dapat dipetakan dengan aman.")
        document.review_blocked = True
    return document


def _parse_legacy_excel(path: Path, manual_mappings: dict[str, dict[str, Any]]) -> ParsedDocument:
    source_hash = file_sha256(path)
    sheets = pd.read_excel(path, sheet_name=None, header=None)
    return _parse_dataframe_sheets(path, source_hash, sheets, "adaptive_xls", manual_mappings)


def _parse_csv(path: Path, manual_mappings: dict[str, dict[str, Any]]) -> ParsedDocument:
    source_hash = file_sha256(path)
    encoding = "utf-8-sig"
    try:
        sample = path.read_text(encoding=encoding)[:8192]
    except UnicodeDecodeError:
        encoding = "latin-1"
        sample = path.read_text(encoding=encoding)[:8192]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    frame = pd.read_csv(path, encoding=encoding, sep=delimiter, header=None, dtype=object)
    return _parse_dataframe_sheets(path, source_hash, {"CSV": frame}, "adaptive_csv", manual_mappings)


def _parse_dataframe_sheets(
    path: Path,
    source_hash: str,
    sheets: dict[str, pd.DataFrame],
    strategy: str,
    manual_mappings: dict[str, dict[str, Any]],
) -> ParsedDocument:
    document = ParsedDocument(path.name, source_hash, strategy)
    sequence = 0
    for sheet_name, frame in sheets.items():
        frame = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
        if frame.empty:
            continue
        header_start, header_end = _find_dataframe_header(frame)
        region_id = f"{sheet_name}:{header_start + 1}-{header_end + 1}"
        headers = _dataframe_header_paths(frame, header_start, header_end)
        data = frame.iloc[header_end + 1 :].reset_index(drop=False)
        mapping = _map_dataframe_columns(headers, data, manual_mappings.get(region_id, {}))
        confidence = _mapping_confidence(mapping)
        blocked = _mapping_blocks_review(mapping, confidence)
        diagnostic = MappingDiagnostic(
            region_id,
            confidence,
            confidence_label(confidence),
            [header_start + 1, header_end + 1] if header_end != header_start else [header_start + 1],
            mapping,
            warnings=_mapping_warnings(mapping),
            review_blocked=blocked,
        )
        items = []
        title = path.stem
        current_section = ""
        for offset, row in data.iterrows():
            desc_col = _mapping_index(mapping, "description")
            if desc_col is None or desc_col >= len(row) - 1:
                continue
            raw_desc = _row_value(row, mapping, "description", dataframe_offset=1)
            desc = _clean_text(raw_desc)
            if not desc:
                continue
            item_no = _clean_item_no(_row_value(row, mapping, "item_number", dataframe_offset=1))
            numeric, raw_values, provenance = _numeric_components_from_dataframe_row(row, mapping)
            if _is_section_row(item_no, desc, numeric):
                current_section = desc
                continue
            if _looks_like_header_or_guide(desc):
                continue
            sequence += 1
            source_row = int(row.iloc[0]) + 1
            item = _build_item(
                path,
                source_hash,
                sheet_name,
                source_row,
                item_no,
                sequence,
                title,
                current_section,
                desc,
                _clean_text(_row_value(row, mapping, "unit", dataframe_offset=1)),
                numeric,
                raw_values,
                strategy,
                confidence,
                provenance,
                coordinate_col=desc_col,
            )
            items.append(item)
        document.items.extend(items)
        document.regions.append(
            TableRegion(region_id, sheet_name, header_start + 1, header_end + 1, header_end + 2, len(frame), diagnostic, strategy)
        )
        if items and blocked:
            document.review_blocked = True
    if not document.items:
        document.review_blocked = True
        document.warnings.append("Tidak ada item RAB yang dapat dipetakan dengan aman.")
    return document


def _meaningful_last_row(value_sheet: Any, formula_sheet: Any) -> int:
    populated_rows = {
        cell.row
        for sheet in (value_sheet, formula_sheet)
        for cell in sheet._cells.values()
        if cell.value not in (None, "")
    }
    return max(populated_rows, default=0)


def _extract_sheet_title(sheet: Any, last_row: int) -> str:
    scan_end = min(last_row, 40)
    for row_no in range(1, scan_end + 1):
        values = [_clean_text(sheet.cell(row_no, col).value) for col in range(1, sheet.max_column + 1)]
        for index, value in enumerate(values):
            normalized = value.lower().replace(":", "").strip()
            if normalized in {"nama pekerjaan", "judul pekerjaan", "pekerjaan"}:
                for candidate in values[index + 1 :]:
                    candidate = candidate.lstrip(": ").strip()
                    if candidate:
                        return candidate
    candidates = []
    for merged in sheet.merged_cells.ranges:
        if merged.min_row > min(scan_end, 12) or merged.max_col - merged.min_col < 3:
            continue
        value = _clean_text(sheet.cell(merged.min_row, merged.min_col).value)
        upper = value.upper()
        if value and len(value) >= 8 and "RAB" not in upper and "PT PLN" not in upper and "UNIT " not in upper:
            candidates.append((merged.min_row, value))
    return sorted(candidates)[-1][1] if candidates else ""


def _find_header_regions(formula_sheet: Any, value_sheet: Any, last_row: int) -> list[tuple[int, int]]:
    max_scan = min(last_row, 250)
    scored: list[tuple[float, int, int]] = []
    for start in range(1, max_scan + 1):
        if formula_sheet.row_dimensions[start].hidden:
            continue
        for depth in range(1, 5):
            end = start + depth - 1
            if end > max_scan:
                break
            headers = [_header_path(formula_sheet, start, end, col) for col in range(1, formula_sheet.max_column + 1)]
            role_best = {role: max((_header_semantic_score(role, header) for header in headers), default=0) for role in ROLE_ORDER}
            numeric_hits = sum(role_best[role] >= 75 for role in NUMERIC_ROLES)
            required = role_best["description"] >= 75 and (
                role_best["unit"] >= 70 or role_best["volume"] >= 70 or numeric_hits >= 2
            )
            if not required:
                continue
            score = role_best["description"] + role_best["item_number"] * 0.35 + numeric_hits * 24 - (depth - 1) * 18
            scored.append((score, start, end))
    selected: list[tuple[int, int]] = []
    for _, start, end in sorted(scored, reverse=True):
        if any(not (end < existing_start - 1 or start > existing_end + 1) for existing_start, existing_end in selected):
            continue
        selected.append((start, end))
    normalized = []
    for start, end in sorted(selected):
        for merged in formula_sheet.merged_cells.ranges:
            if merged.min_row <= end and merged.max_row >= start:
                value = _clean_text(formula_sheet.cell(merged.min_row, merged.min_col).value)
                if max(_header_semantic_score(role, value) for role in ROLE_ORDER) >= 75:
                    start = min(start, merged.min_row)
        for row_no in range(start, min(last_row, end + 2) + 1):
            if _is_header_guide_row(formula_sheet, row_no):
                end = row_no - 1
                break
        candidate = (start, max(start, end))
        if candidate not in normalized:
            normalized.append(candidate)
    return sorted(normalized)


def _is_header_guide_row(sheet: Any, row_no: int) -> bool:
    values = [
        _clean_text(sheet.cell(row_no, col).value)
        for col in range(1, sheet.max_column + 1)
        if _clean_text(sheet.cell(row_no, col).value)
    ]
    if len(values) < 3:
        return False
    guide_values = sum(
        bool(re.fullmatch(r"\d+(?:=\d+(?:x|\+)\d+)?", value.replace(" ", "").lower()))
        for value in values
    )
    return guide_values >= max(3, math.ceil(len(values) * 0.6))


def _header_path(sheet: Any, start: int, end: int, col: int) -> str:
    values = []
    for row in range(start, end + 1):
        value = _merged_value(sheet, row, col)
        text = _clean_text(value)
        if text and text not in values:
            values.append(text)
    return " | ".join(values)


def _merged_value(sheet: Any, row: int, col: int) -> Any:
    cache_key = (id(sheet), row, col)
    if cache_key in _MERGED_VALUE_CACHE:
        return _MERGED_VALUE_CACHE[cache_key]
    cell = sheet.cell(row, col)
    if cell.value not in (None, ""):
        _MERGED_VALUE_CACHE[cache_key] = cell.value
        return cell.value
    coordinate = cell.coordinate
    for merged in sheet.merged_cells.ranges:
        if coordinate in merged:
            value = sheet.cell(merged.min_row, merged.min_col).value
            _MERGED_VALUE_CACHE[cache_key] = value
            return value
    _MERGED_VALUE_CACHE[cache_key] = cell.value
    return cell.value


def _map_columns(
    formula_sheet: Any,
    value_sheet: Any,
    header_start: int,
    header_end: int,
    data_end: int,
    manual: dict[str, Any],
) -> dict[str, ColumnMapping]:
    mappings = {}
    header_columns = [
        col
        for col in range(1, max(formula_sheet.max_column, value_sheet.max_column) + 1)
        if _header_path(formula_sheet, header_start, header_end, col)
    ]
    max_col = max(header_columns, default=max(formula_sheet.max_column, value_sheet.max_column))
    arithmetic_evidence = _column_arithmetic_evidence(
        formula_sheet, header_end + 1, data_end, max_col
    )
    for role in ROLE_ORDER:
        candidates = []
        for col in range(1, max_col + 1):
            header = _header_path(formula_sheet, header_start, header_end, col)
            values = [
                value_sheet.cell(row, col).value
                for row in range(header_end + 1, min(data_end, header_end + 40) + 1)
                if not formula_sheet.row_dimensions[row].hidden
            ]
            formulas = [
                formula_sheet.cell(row, col).value
                for row in range(header_end + 1, min(data_end, header_end + 40) + 1)
                if not formula_sheet.row_dimensions[row].hidden
            ]
            header_score = _header_semantic_score(role, header)
            if header_score and _header_is_only_inherited(formula_sheet, header_start, header_end, col):
                header_score *= 0.35
            value_score = _value_pattern_score(role, values)
            arithmetic_score = _arithmetic_score(role, col, arithmetic_evidence)
            structure_score = _structure_score(role, col, max_col)
            total = 0.45 * header_score + 0.25 * value_score + 0.20 * arithmetic_score + 0.10 * structure_score
            candidates.append((total, col, header, header_score, value_score, arithmetic_score, structure_score, formulas))
        candidates.sort(reverse=True, key=lambda item: item[0])
        winner = candidates[0]
        runner_up = candidates[1][0] if len(candidates) > 1 else 0.0
        manual_col = _manual_column_index(manual.get(role), max_col)
        confirmed = manual_col is not None
        if confirmed:
            winner = next((item for item in candidates if item[1] == manual_col), winner)
            runner_up = 0.0
        accepted = confirmed or (winner[0] >= 65 and winner[0] - runner_up >= 8)
        mappings[role] = ColumnMapping(
            role=role,
            column_index=winner[1] if accepted else None,
            column_label=_column_letter(winner[1]) if accepted else "",
            header_text=winner[2],
            score=round(100.0 if confirmed else winner[0], 2),
            header_score=round(winner[3], 2),
            value_score=round(winner[4], 2),
            arithmetic_score=round(winner[5], 2),
            structure_score=round(winner[6], 2),
            runner_up_score=round(runner_up, 2),
            ambiguous=not accepted and winner[0] >= 50,
            confirmed_manually=confirmed,
        )
    return mappings


def _header_is_only_inherited(sheet: Any, header_start: int, header_end: int, col: int) -> bool:
    if any(sheet.cell(row, col).value not in (None, "") for row in range(header_start, header_end + 1)):
        return False
    return any(
        merged.min_col < col <= merged.max_col
        and merged.min_row <= header_end
        and merged.max_row >= header_start
        for merged in sheet.merged_cells.ranges
    )


def _map_dataframe_columns(headers: list[str], data: pd.DataFrame, manual: dict[str, Any]) -> dict[str, ColumnMapping]:
    mappings = {}
    max_col = len(headers)
    for role in ROLE_ORDER:
        candidates = []
        for col, header in enumerate(headers, start=1):
            values = data.iloc[:, col].tolist() if col < len(data.columns) else []
            hs = _header_semantic_score(role, header)
            vs = _value_pattern_score(role, values)
            ars = 50.0 if role in NUMERIC_ROLES and vs >= 60 else 20.0
            ss = _structure_score(role, col, max_col)
            score = 0.45 * hs + 0.25 * vs + 0.20 * ars + 0.10 * ss
            candidates.append((score, col, header, hs, vs, ars, ss))
        candidates.sort(reverse=True)
        winner = candidates[0]
        runner = candidates[1][0] if len(candidates) > 1 else 0.0
        manual_col = _manual_column_index(manual.get(role), max_col)
        confirmed = manual_col is not None
        if confirmed:
            winner = next((item for item in candidates if item[1] == manual_col), winner)
            runner = 0.0
        accepted = confirmed or (winner[0] >= 65 and winner[0] - runner >= 8)
        mappings[role] = ColumnMapping(
            role,
            winner[1] if accepted else None,
            str(winner[1]) if accepted else "",
            winner[2],
            round(100.0 if confirmed else winner[0], 2),
            round(winner[3], 2),
            round(winner[4], 2),
            round(winner[5], 2),
            round(winner[6], 2),
            round(runner, 2),
            not accepted and winner[0] >= 50,
            confirmed,
        )
    return mappings


def _header_semantic_score(role: str, header: str) -> float:
    text = re.sub(r"\s+", " ", str(header or "").lower()).strip()
    if not text:
        return 0.0
    has_material = any(token in text for token in ("material", "barang"))
    has_service = any(token in text for token in ("jasa", "service"))
    has_currency = any(token in text for token in ("rp", "harga", "price", "satuan"))
    has_amount = any(token in text for token in ("jumlah", "total", "nilai", "amount"))
    if role == "item_number":
        return 100.0 if re.search(r"(?:^|\|)\s*(?:no\.?|nomor|urut)\b", text) else 0.0
    if role == "description":
        return 100.0 if any(token in text for token in ("nama barang", "nama material", "uraian", "deskripsi", "pekerjaan", "keterangan item")) else 0.0
    if role == "unit":
        if any(token in text for token in ("harga", "jumlah", "total")):
            return 0.0
        return 100.0 if re.search(r"(?:^|\|)\s*(?:sat\.?|satuan|unit|uom)\s*(?:\||$)", text) else 0.0
    if role == "volume":
        return 100.0 if re.search(r"\b(?:vol\.?|volume|qty|kuantitas)\b", text) else 0.0
    if role == "material_unit_price":
        return 100.0 if has_material and has_currency and not has_amount else 0.0
    if role == "service_unit_price":
        return 100.0 if has_service and has_currency and not has_amount else 0.0
    if role == "unit_price":
        return 95.0 if ("harga satuan" in text or "unit price" in text) and not (has_material or has_service) else 0.0
    if role == "material_total":
        return 100.0 if has_material and has_amount else 0.0
    if role == "service_total":
        return 100.0 if has_service and has_amount else 0.0
    if role == "total_price":
        if has_material or has_service:
            return 0.0
        return 100.0 if re.search(r"(?:^|\|)\s*(?:total|jumlah|nilai)(?:\s|\||$)", text) else 0.0
    if role == "notes":
        return 100.0 if any(token in text for token in ("catatan", "remark", "notes", "keterangan")) else 0.0
    return 0.0


def _value_pattern_score(role: str, values: Iterable[Any]) -> float:
    usable = [value for value in values if _clean_text(value).lower() not in MARKERS]
    if not usable:
        return 0.0
    numeric_ratio = sum(parse_number(value) is not None for value in usable) / len(usable)
    text_ratio = sum(parse_number(value) is None and bool(_clean_text(value)) for value in usable) / len(usable)
    if role in NUMERIC_ROLES:
        return numeric_ratio * 100
    if role == "description":
        descriptive = sum(len(_clean_text(value)) >= 5 and parse_number(value) is None for value in usable)
        return descriptive / len(usable) * 100
    if role == "unit":
        text_values = [value for value in usable if parse_number(value) is None]
        if not text_values:
            return 0.0
        short_text = sum(0 < len(_clean_text(value)) <= 12 for value in text_values)
        return short_text / len(text_values) * 100
    if role == "item_number":
        identifiers = sum(bool(re.fullmatch(r"[0-9]+(?:\.[0-9]+)*|[A-Za-z]{1,4}|[IVXLCDM]+", _clean_text(value))) for value in usable)
        return identifiers / len(usable) * 100
    return text_ratio * 100


def _column_arithmetic_evidence(sheet: Any, data_start: int, data_end: int, max_col: int) -> dict[str, Any]:
    own_formulas = {col: 0 for col in range(1, max_col + 1)}
    references = {col: 0 for col in range(1, max_col + 1)}
    formula_count = 0
    for row in range(data_start, min(data_end, data_start + 50) + 1):
        for col in range(1, max_col + 1):
            formula = sheet.cell(row, col).value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            formula_count += 1
            own_formulas[col] += 1
            for column_letters, referenced_row in re.findall(r"\$?([A-Z]{1,3})\$?(\d+)", formula.upper()):
                if int(referenced_row) != row:
                    continue
                referenced_col = _manual_column_index(column_letters, max_col)
                if referenced_col:
                    references[referenced_col] += 1
    return {"own": own_formulas, "references": references, "formula_count": formula_count}


def _arithmetic_score(role: str, col: int, evidence: dict[str, Any]) -> float:
    references = evidence["references"].get(col, 0)
    own_formulas = evidence["own"].get(col, 0)
    if role in {"volume", "material_unit_price", "service_unit_price", "unit_price"}:
        return min(100.0, references * 22.0)
    if role in {"material_total", "service_total", "total_price"}:
        return min(100.0, own_formulas * 14.0 + references * 8.0)
    if role == "description":
        return 50.0
    return 20.0 if evidence["formula_count"] else 0.0


def _structure_score(role: str, col: int, max_col: int) -> float:
    position = col / max(max_col, 1)
    expected = {
        "item_number": 0.05,
        "description": 0.18,
        "unit": 0.32,
        "volume": 0.55,
        "material_unit_price": 0.67,
        "service_unit_price": 0.72,
        "unit_price": 0.68,
        "material_total": 0.80,
        "service_total": 0.86,
        "total_price": 0.95,
        "notes": 0.95,
    }.get(role, 0.5)
    return max(0.0, 100.0 - abs(position - expected) * 180.0)


def _mapping_confidence(mapping: dict[str, ColumnMapping]) -> float:
    required_roles = ["description", "unit", "volume", "total_price"]
    component_roles = ["material_unit_price", "service_unit_price", "unit_price", "material_total", "service_total"]
    scores = [mapping[role].score for role in required_roles if mapping[role].column_index is not None]
    scores.extend(sorted((mapping[role].score for role in component_roles if mapping[role].column_index is not None), reverse=True)[:2])
    if mapping["description"].column_index is None:
        return 0.0
    if not scores:
        return 0.0
    coverage = min(1.0, len(scores) / 5.0)
    return round(sum(scores) / len(scores) * (0.75 + 0.25 * coverage), 2)


def _mapping_warnings(mapping: dict[str, ColumnMapping]) -> list[str]:
    warnings = []
    for role, item in mapping.items():
        if item.ambiguous:
            warnings.append(f"Mapping {role} ambigu; konfirmasi kolom secara manual.")
    if mapping["description"].column_index is None:
        warnings.append("Kolom uraian belum dapat ditentukan.")
    if not any(mapping[role].column_index is not None for role in NUMERIC_ROLES):
        warnings.append("Komponen angka RAB belum dapat ditentukan.")
    return warnings


def _mapping_blocks_review(mapping: dict[str, ColumnMapping], confidence: float) -> bool:
    manual_confirmed = any(item.confirmed_manually for item in mapping.values())
    essential_ambiguous = mapping["description"].ambiguous
    numeric_mapped = any(mapping[role].column_index is not None for role in NUMERIC_ROLES)
    return not manual_confirmed and (
        confidence < 65
        or mapping["description"].column_index is None
        or essential_ambiguous
        or not numeric_mapped
    )


def _extract_region_items(
    path: Path,
    source_hash: str,
    formula_sheet: Any,
    value_sheet: Any,
    title: str,
    data_start: int,
    data_end: int,
    mapping: dict[str, ColumnMapping],
    parser_confidence: float,
    display_sequence: int,
) -> tuple[list[RABItem], int, int]:
    items = []
    failures = 0
    major_section = ""
    minor_section = ""
    for row_no in range(data_start, data_end + 1):
        if formula_sheet.row_dimensions[row_no].hidden:
            continue
        desc = _clean_text(_sheet_value(value_sheet, row_no, mapping, "description"))
        if not desc or _looks_like_header_or_guide(desc):
            continue
        item_no = _clean_item_no(_sheet_value(value_sheet, row_no, mapping, "item_number"))
        numeric, raw_values, provenance = _numeric_components_from_sheet(
            formula_sheet, value_sheet, row_no, mapping
        )
        if _is_section_row(item_no, desc, numeric):
            if re.fullmatch(r"[IVXLCDM]+", item_no.upper()):
                major_section, minor_section = desc, ""
            else:
                minor_section = desc
            continue
        if not item_no and not any(value is not None for value in numeric.values()):
            continue
        display_sequence += 1
        section = " > ".join(part for part in (major_section, minor_section) if part)
        warnings = _arithmetic_warnings(numeric)
        if warnings:
            failures += 1
        unit = _clean_text(_sheet_value(value_sheet, row_no, mapping, "unit"))
        notes = _clean_text(_sheet_value(value_sheet, row_no, mapping, "notes"))
        item = _build_item(
            path,
            source_hash,
            formula_sheet.title,
            row_no,
            item_no,
            display_sequence,
            title,
            section,
            desc,
            unit,
            numeric,
            raw_values,
            "adaptive_xlsx",
            parser_confidence,
            provenance,
            warnings,
            notes,
            coordinate_col=_mapping_index(mapping, "description") or 1,
        )
        items.append(item)
    return items, failures, display_sequence


def _numeric_components_from_sheet(
    formula_sheet: Any,
    value_sheet: Any,
    row_no: int,
    mapping: dict[str, ColumnMapping],
) -> tuple[dict[str, float | None], dict[str, Any], dict[str, str]]:
    values: dict[str, float | None] = {}
    raw: dict[str, Any] = {}
    provenance: dict[str, str] = {}
    for role in NUMERIC_ROLES:
        col = _mapping_index(mapping, role)
        cached = value_sheet.cell(row_no, col).value if col else None
        formula = formula_sheet.cell(row_no, col).value if col else None
        raw[role] = {"value": cached, "formula": formula if isinstance(formula, str) and formula.startswith("=") else ""}
        values[role] = parse_number(cached)
        provenance[role] = "source" if values[role] is not None else "missing"
    _derive_financial_values(values, provenance)
    return values, raw, provenance


def _numeric_components_from_dataframe_row(
    row: pd.Series,
    mapping: dict[str, ColumnMapping],
) -> tuple[dict[str, float | None], dict[str, Any], dict[str, str]]:
    raw = {
        role: {"value": _row_value(row, mapping, role, dataframe_offset=1), "formula": ""}
        for role in NUMERIC_ROLES
    }
    values = {role: parse_number(raw[role]["value"]) for role in NUMERIC_ROLES}
    provenance = {role: "source" if value is not None else "missing" for role, value in values.items()}
    _derive_financial_values(values, provenance)
    return values, raw, provenance


def _derive_financial_values(values: dict[str, float | None], provenance: dict[str, str]) -> None:
    if values.get("unit_price") is None:
        components = [values.get("material_unit_price"), values.get("service_unit_price")]
        available = [value for value in components if value is not None]
        if available:
            values["unit_price"] = sum(available)
            provenance["unit_price"] = "derived"
    if values.get("total_price") is None:
        components = [values.get("material_total"), values.get("service_total")]
        available = [value for value in components if value is not None]
        if available:
            values["total_price"] = sum(available)
            provenance["total_price"] = "derived"
        elif values.get("volume") is not None and values.get("unit_price") is not None:
            values["total_price"] = float(values["volume"]) * float(values["unit_price"])
            provenance["total_price"] = "derived"


def _arithmetic_warnings(values: dict[str, float | None]) -> list[str]:
    warnings = []
    volume = values.get("volume")
    unit_price = values.get("unit_price")
    total = values.get("total_price")
    if volume is not None and unit_price is not None and total is not None:
        expected = volume * unit_price
        if not _approximately_equal(total, expected):
            warnings.append(f"Total sumber {total:g} tidak sama dengan volume x harga satuan {expected:g}.")
    component_totals = [values.get("material_total"), values.get("service_total")]
    available = [value for value in component_totals if value is not None]
    if total is not None and available and not _approximately_equal(total, sum(available)):
        warnings.append(f"Total sumber {total:g} tidak sama dengan jumlah komponen {sum(available):g}.")
    return warnings


def _approximately_equal(left: float, right: float) -> bool:
    return abs(left - right) <= max(1.0, abs(right) * 0.005)


def _build_item(
    path: Path,
    source_hash: str,
    sheet: str,
    source_row: int,
    item_no: str,
    sequence: int,
    title: str,
    section: str,
    description: str,
    unit: str,
    numeric: dict[str, Any],
    raw_values: dict[str, Any],
    strategy: str,
    confidence: float,
    provenance: dict[str, str] | None = None,
    warnings: list[str] | None = None,
    notes: str = "",
    coordinate_col: int = 1,
) -> RABItem:
    coordinate = f"{_column_letter(coordinate_col)}{source_row}"
    source_id = f"{source_hash[:12]}:{sheet}!{coordinate}"
    provenance = provenance or numeric.pop("_provenance", {})
    return RABItem(
        source_id=source_id,
        source_file=path.name,
        source_hash=source_hash,
        source_location=f"{sheet}!{coordinate}",
        page_or_sheet=sheet,
        source_row=source_row,
        source_coordinate=coordinate,
        item_no=item_no,
        display_sequence=sequence,
        judul_rab=title,
        section=section,
        item_per_rab=description,
        unit=unit,
        volume=numeric.get("volume"),
        material_unit_price=numeric.get("material_unit_price"),
        service_unit_price=numeric.get("service_unit_price"),
        unit_price=numeric.get("unit_price"),
        material_total=numeric.get("material_total"),
        service_total=numeric.get("service_total"),
        total_price=numeric.get("total_price"),
        notes=notes,
        raw_values=raw_values,
        provenance=provenance,
        parser_strategy=strategy,
        parser_confidence=confidence,
        parser_warnings=warnings or [],
    )


def _find_dataframe_header(frame: pd.DataFrame) -> tuple[int, int]:
    best = (-1.0, 0, 0)
    scan = min(len(frame), 80)
    for start in range(scan):
        for depth in range(1, min(4, scan - start) + 1):
            end = start + depth - 1
            paths = _dataframe_header_paths(frame, start, end)
            desc = max(_header_semantic_score("description", path) for path in paths)
            roles = sum(max(_header_semantic_score(role, path) for path in paths) >= 75 for role in NUMERIC_ROLES | {"unit", "volume"})
            score = desc + roles * 20 - (depth - 1) * 18
            if desc >= 75 and score > best[0]:
                best = (score, start, end)
    return (best[1], best[2]) if best[0] >= 0 else (0, 0)


def _dataframe_header_paths(frame: pd.DataFrame, start: int, end: int) -> list[str]:
    paths = []
    for col in range(frame.shape[1]):
        values = []
        for row in range(start, end + 1):
            text = _clean_text(frame.iat[row, col])
            if text and text not in values:
                values.append(text)
        paths.append(" | ".join(values))
    return paths


def _sheet_value(sheet: Any, row: int, mapping: dict[str, ColumnMapping], role: str) -> Any:
    col = _mapping_index(mapping, role)
    return sheet.cell(row, col).value if col else None


def _row_value(row: pd.Series, mapping: dict[str, ColumnMapping], role: str, dataframe_offset: int = 0) -> Any:
    col = _mapping_index(mapping, role)
    if col is None:
        return None
    position = col - 1 + dataframe_offset
    return row.iloc[position] if 0 <= position < len(row) else None


def _mapping_index(mapping: dict[str, ColumnMapping], role: str) -> int | None:
    item = mapping.get(role)
    return item.column_index if item else None


def _manual_column_index(value: Any, max_col: int) -> int | None:
    if value in (None, "", 0, "0"):
        return None
    if isinstance(value, int):
        return value if 1 <= value <= max_col else None
    text = str(value).strip().upper()
    if text.isdigit():
        number = int(text)
        return number if 1 <= number <= max_col else None
    number = 0
    for char in text:
        if not ("A" <= char <= "Z"):
            return None
        number = number * 26 + ord(char) - 64
    return number if 1 <= number <= max_col else None


def _column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_item_no(value: Any) -> str:
    text = _clean_text(value)
    return text[:-2] if text.endswith(".0") else text


def _is_section_row(item_no: str, description: str, numeric: dict[str, Any]) -> bool:
    has_numbers = any(numeric.get(role) is not None for role in NUMERIC_ROLES)
    section_id = bool(re.fullmatch(r"[A-Za-z]{1,4}|[IVXLCDM]+", item_no or "", re.IGNORECASE))
    return bool(description and not has_numbers and section_id)


def _looks_like_header_or_guide(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", str(text or "").lower()).strip()
    if not normalized:
        return True
    if re.fullmatch(r"\d+(?:\s*=.*)?", normalized):
        return True
    tokens = ("nama barang", "nama material", "uraian pekerjaan", "harga satuan", "rencana anggaran dan biaya")
    return any(token in normalized for token in tokens)
