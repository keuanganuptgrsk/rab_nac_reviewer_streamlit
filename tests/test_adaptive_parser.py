from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from modules.adaptive_parser import parse_document, parse_number


def _merged_rab_fixture(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RAB"
    sheet.merge_cells("A4:P4")
    sheet["A4"] = "RINCIAN RENCANA ANGGARAN DAN BIAYA (RAB)"
    sheet.merge_cells("A5:P5")
    sheet["A5"] = "RENOVASI AREA KERJA"
    sheet.merge_cells("A8:A10")
    sheet["A8"] = "No."
    sheet.merge_cells("B8:D10")
    sheet["B8"] = "Nama Barang ( Material ) dan Jasa"
    sheet.merge_cells("E8:P8")
    sheet["E8"] = "RENCANA ANGGARAN DAN BIAYA (RAB)"
    sheet.merge_cells("E9:E10")
    sheet["E9"] = "Sat."
    sheet.merge_cells("K9:K10")
    sheet["K9"] = "Vol."
    sheet.merge_cells("L9:M9")
    sheet["L9"] = "SATUAN"
    sheet["L10"] = "Material (Rp.)"
    sheet["M10"] = "Jasa (Rp.)"
    sheet.merge_cells("N9:O9")
    sheet["N9"] = "JUMLAH"
    sheet["N10"] = "Material (Rp.)"
    sheet["O10"] = "Jasa (Rp.)"
    sheet["P9"] = "TOTAL"
    sheet["P10"] = "(Rp.)"
    for coordinate, value in {
        "A11": 1,
        "B11": 2,
        "E11": 3,
        "K11": 4,
        "L11": 5,
        "M11": 6,
        "N11": "7=4x5",
        "O11": "8=4x6",
        "P11": "9=7+8",
        "A12": "A",
        "B12": "PEKERJAAN PERSIAPAN",
        "A13": 1,
        "B13": "Mobilisasi dan Demobilisasi",
        "E13": "ls",
        "K13": 1,
        "L13": "include",
        "M13": 2136100,
        "N13": "include",
        "O13": 2136100,
        "P13": 2136100,
        "A14": 2,
        "B14": "Pembongkaran keramik",
        "E14": "m2",
        "K14": 17,
        "L14": "include",
        "M14": 27900,
        "N14": "include",
        "O14": 474300,
        "P14": 474300,
        "A15": 3,
        "B15": "Baris tersembunyi",
        "E15": "ls",
        "K15": 1,
        "M15": 999999,
        "O15": 999999,
        "P15": 999999,
    }.items():
        sheet[coordinate] = value
    sheet.row_dimensions[15].hidden = True
    workbook.save(path)
    return path


def test_merged_header_fixture_maps_exact_financial_values(tmp_path):
    document = parse_document(_merged_rab_fixture(tmp_path / "merged.xlsx"))

    assert document.review_blocked is False
    assert len(document.items) == 2
    first, second = document.items
    assert (first.unit, first.volume, first.unit_price, first.total_price) == ("ls", 1, 2136100, 2136100)
    assert (second.unit, second.volume, second.unit_price, second.total_price) == ("m2", 17, 27900, 474300)
    assert first.source_id.endswith("RAB!B13")
    assert first.provenance["unit_price"] == "derived"
    assert document.regions[0].diagnostic.hidden_rows_skipped == 1


def test_service_only_and_material_only_are_combined_safely(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["No", "Uraian", "Satuan", "Volume", "Harga Material", "Harga Jasa", "Total"])
    sheet.append([1, "Material panel", "unit", 2, 100000, "N/A", 200000])
    sheet.append([2, "Jasa instalasi", "ls", 1, "N/A", 350000, 350000])
    path = tmp_path / "components.xlsx"
    workbook.save(path)

    document = parse_document(path)

    assert [item.unit_price for item in document.items] == [100000, 350000]
    assert [item.total_price for item in document.items] == [200000, 350000]


def test_csv_flat_strategy_and_indonesian_numbers(tmp_path):
    path = tmp_path / "rab.csv"
    path.write_text("uraian;volume;satuan;total\nKonsumsi rapat;2;paket;Rp 1.250.000\n", encoding="utf-8")

    document = parse_document(path)

    assert len(document.items) == 1
    assert document.items[0].total_price == 1250000
    assert document.items[0].source_id.endswith("CSV!A2")
    assert document.items[0].raw_values["total_price"]["value"] == "Rp 1.250.000"


def test_number_parser_preserves_markers_as_unknown():
    assert parse_number("Rp 1.234.567,89") == 1234567.89
    assert parse_number("1,234.50") == 1234.5
    assert parse_number("include") is None
    assert parse_number("N/A") is None
    assert parse_number("-") is None
    assert parse_number("Item pekerjaan nomor 1") is None


def test_simple_dataframe_does_not_use_data_row_as_title(tmp_path):
    path = tmp_path / "flat.xlsx"
    pd.DataFrame({"uraian": ["Item pertama"], "volume": [1], "total": [20000]}).to_excel(path, index=False)

    document = parse_document(path)

    assert document.items[0].judul_rab == "flat"


def test_ambiguous_description_is_blocked_until_manual_mapping(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ambigu"
    sheet.append(["Uraian", "Uraian", "Volume", "Total"])
    sheet.append(["Konsumsi rapat", "Catatan item", 1, 500000])
    path = tmp_path / "ambiguous.xlsx"
    workbook.save(path)

    blocked = parse_document(path)
    assert blocked.review_blocked is True

    confirmed = parse_document(
        path,
        {"Ambigu:1-1": {"description": "A", "volume": "C", "total_price": "D"}},
    )
    assert confirmed.review_blocked is False
    assert confirmed.items[0].item_per_rab == "Konsumsi rapat"
    assert confirmed.items[0].total_price == 500000
