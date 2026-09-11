from __future__ import annotations

import importlib
import sqlite3
import time
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook


def _load_modules(monkeypatch, tmp_path):
    monkeypatch.setenv("RAB_NAC_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("RAB_NAC_DB_PATH", str(tmp_path / "data" / "app.db"))
    monkeypatch.setenv("RAB_NAC_EXPORT_DIR", str(tmp_path / "exports"))
    import modules.db as db
    import modules.export_engine as export_engine

    importlib.reload(db)
    importlib.reload(export_engine)
    db.init_db()
    return db, export_engine


def _long_result():
    long_reason = " ".join(f"bukti-audit-{index}" for index in range(90))
    return {
        "source_id": "abc123:RAB!B13",
        "row_id": "1",
        "source_file": "rab.xlsx",
        "page_or_sheet": "RAB",
        "source_coordinate": "B13",
        "judul_rab": "Renovasi ruang kerja",
        "section": "Pekerjaan pendukung",
        "item_per_rab": "Renovasi ruang kerja dan fasilitas pendukung untuk unit finance",
        "unit": "ls",
        "volume": 1,
        "unit_price": 1000000,
        "total_price": 1000000,
        "matched_keyword": "Renovasi Ruang Kerja",
        "matched_category": "Kategori B - Koreksi BPP",
        "selected_transaction_type": "Renovasi ruang kerja/pendukung",
        "reference_percentage": 20,
        "applied_nac_percentage": 20,
        "correction_percentage_label": "20%",
        "percentage_status": "Diterapkan dari transaksi terpilih",
        "final_confidence": 73.5,
        "deterministic_confidence": 73.5,
        "confidence_label": "Tinggi",
        "recommended_action": "Perlu Review Manual",
        "decision_reason": long_reason,
        "explanation": long_reason,
        "engine": "Python Lokal",
        "provider_model": "deterministic-nac-engine",
        "decision_source": "deterministic_local",
        "parser_strategy": "adaptive_xlsx",
        "parser_confidence": 91.2,
        "provenance": {"total_price": "source"},
    }


def test_pdf_preserves_long_audit_text_and_has_page_numbers(monkeypatch, tmp_path):
    _, export_engine = _load_modules(monkeypatch, tmp_path)
    empty_percentage = _long_result()
    empty_percentage["item_per_rab"] = "Item tanpa prosentase"
    empty_percentage["applied_nac_percentage"] = float("nan")
    path = Path(export_engine.export_potential_nac_pdf([_long_result(), empty_percentage]))

    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
        assert document.page_count >= 2
        assert "bukti-audit-89" in text
        assert "nan%" not in text.lower()
        assert "Internal Review Draft" in text
        assert f"Page {document.page_count} of {document.page_count}" in text


def test_excel_export_has_finance_formatting_and_numeric_values(monkeypatch, tmp_path):
    _, export_engine = _load_modules(monkeypatch, tmp_path)
    path = Path(export_engine.export_all_materials_excel([_long_result()]))
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Seluruh Material"]

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert isinstance(sheet.cell(2, headers["Total"]).value, (int, float))
    assert "#,##0" in sheet.cell(2, headers["Total"]).number_format
    assert sheet.cell(1, 1).fill.fgColor.rgb.endswith("246B61")


def test_invalid_sqlite_restore_is_rejected_without_replacing_database(monkeypatch, tmp_path):
    db, _ = _load_modules(monkeypatch, tmp_path)
    original_count = len(db.get_keywords(False))
    invalid = tmp_path / "invalid.db"
    invalid.write_bytes(b"not a sqlite database")

    with pytest.raises(ValueError):
        db.restore_db(invalid)

    assert len(db.get_keywords(False)) == original_count


def test_sqlite_restore_requires_application_schema(monkeypatch, tmp_path):
    db, _ = _load_modules(monkeypatch, tmp_path)
    foreign = tmp_path / "foreign.db"
    with sqlite3.connect(foreign) as connection:
        connection.execute("CREATE TABLE unrelated (id INTEGER PRIMARY KEY)")

    with pytest.raises(ValueError, match="Tabel wajib"):
        db.restore_db(foreign)


def test_restore_rolls_back_when_initialization_fails(monkeypatch, tmp_path):
    db, _ = _load_modules(monkeypatch, tmp_path)
    source = tmp_path / "valid-backup.db"
    db.backup_db(source)
    original = db.DB_PATH.read_bytes()

    def fail_init():
        raise RuntimeError("simulated migration failure")

    monkeypatch.setattr(db, "init_db", fail_init)
    with pytest.raises(RuntimeError, match="migration failure"):
        db.restore_db(source)

    assert db.DB_PATH.read_bytes() == original


def test_upload_validation_sanitizes_path_and_rejects_invalid_content(monkeypatch, tmp_path):
    monkeypatch.setenv("RAB_NAC_UPLOAD_DIR", str(tmp_path / "uploads"))
    import modules.review_flow as review_flow

    importlib.reload(review_flow)

    class Upload:
        def __init__(self, name, data):
            self.name = name
            self._data = data

        def getbuffer(self):
            return self._data

    saved = review_flow.save_uploaded_file(
        Upload("../../RAB finance.xlsx", b"PK\x03\x04sample"),
        session_id="session/finance",
    )
    assert saved.parent.name == "session_finance"
    assert saved.name == "RAB_finance.xlsx"
    assert saved.is_relative_to(tmp_path / "uploads")

    with pytest.raises(ValueError, match="tidak sesuai"):
        review_flow.save_uploaded_file(Upload("scan.pdf", b"not-pdf"))

    monkeypatch.setattr(review_flow, "MAX_UPLOAD_BYTES", 4)
    with pytest.raises(ValueError, match="batas"):
        review_flow.save_uploaded_file(Upload("rab.csv", b"12345"))


def test_stale_upload_directories_are_cleaned(monkeypatch, tmp_path):
    import modules.review_flow as review_flow

    stale = tmp_path / "uploads" / "stale-session"
    stale.mkdir(parents=True)
    old = time.time() - 100
    import os

    os.utime(stale, (old, old))
    review_flow.cleanup_stale_uploads(tmp_path / "uploads", ttl_seconds=10)

    assert not stale.exists()
